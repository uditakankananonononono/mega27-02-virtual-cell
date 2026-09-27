"""Fixed 14-ID lookup in published, reused S4/S7; not independent validation."""
import hashlib,json
from pathlib import Path
import openpyxl
root=Path('/tmp/nogales-supp');s4=root/'EMI-22-255-s007.xlsx';s7=root/'EMI-22-255-s010.xlsx';s1=root/'EMI-22-255-s001.docx'
assert hashlib.sha256(s4.read_bytes()).hexdigest()=='73afee37058fab4ef254f3d0acd83fd9eeac9f48939d832a7f94c78b5bd7e7ab'
assert hashlib.sha256(s7.read_bytes()).hexdigest()=='3788cb2009b295c54301fa939cb28d107e317b90f1835737d465860b49cca322'
mo=json.load(open('results/moco_subset_interaction_ijn1463.json'));mo_ids=sorted(next(z for z in mo['subsets'] if z['removed']==['bmocogdp_c','mocogdp_c'])['rescued_ko_growth'])
ex=json.load(open('results/postresult_pair_exception_reactions.json'));pair_ids=sorted(next(z for z in ex['exceptions'] if z['pair']==['2fe2s_c','btamp_c'])['lost_union_genes'])
assert len(mo_ids)==6 and len(pair_ids)==8 and len(set(mo_ids+pair_ids))==14
ws=openpyxl.load_workbook(s4,read_only=True,data_only=True).active;columns={2:'PRCC Genes',3:'Model Genes',4:'PRCC Genes in Model',5:'iLB Essential Genes',6:'iLB Genes in PRCC',7:'Glucose Essential Genes',18:'unlabeled left comparison panel Gene ID',24:'unlabeled right comparison panel Gene ID'}
mem={g:[] for g in mo_ids+pair_ids}
for row in ws.iter_rows(values_only=True):
 for col,label in columns.items():
  if row[col-1] in mem:mem[row[col-1]].append(label)
ws7=openpyxl.load_workbook(s7,read_only=True,data_only=True).active;rows={r[0]:r for r in list(ws7.values)[1:] if r[0]}
out={'status':'reused publisher validation table target lookup; not independent phenotype or gate','protocol':'notes/postresult_reused_phenotype_contradiction_protocol.md','source_url':'https://www.ebi.ac.uk/europepmc/webservices/rest/PMC7078882/supplementaryFiles','s4_sha256':hashlib.sha256(s4.read_bytes()).hexdigest(),'s7_sha256':hashlib.sha256(s7.read_bytes()).hexdigest(),'s1_sha256':hashlib.sha256(s1.read_bytes()).hexdigest(),'publisher_s7_threshold':-2.7,'panels':{},'limits':'S7 is already-published model validation with a cutoff selected on the data and checked against PRCC; S7 fitness is a competition metric, not direct knockout survival. S4 has model-only and unlabeled comparison columns; no phenotype inferred from mere gene membership. Missing S7 target rows are unknown, not negative. Different model versions, medium choices and editing interventions prevent direct biological validation.'}
for panel,ids in [('MoCo',mo_ids),('suppressive_pair',pair_ids)]:
 out['panels'][panel]={}
 for gene in ids:
  r=rows.get(gene)
  z={'s4_memberships':mem[gene], 's7':None}
  if r:
   z['s7']={k:{'fitness':float(r[i]),'published_result_code':r[i+1],'fitness_above_minus_2point7':float(r[i])>-2.7} for k,i in [('glucose',1),('acetate',3),('p_coumarate',5)]}
  out['panels'][panel][gene]=z
all_found=[z['s7'] for panel in out['panels'].values() for z in panel.values() if z['s7']]
assert len(all_found)==11 and all(y['published_result_code']=='FN' and y['fitness_above_minus_2point7'] for x in all_found for y in x.values())
Path('results/postresult_reused_phenotype_targets.json').write_text(json.dumps(out,indent=2,sort_keys=True)+'\n')
print('S7 present',len(all_found),'all FN each substrate:',all(y['published_result_code']=='FN' for x in all_found for y in x.values()))
for panel,genes in out['panels'].items():
 print(panel)
 for g,z in genes.items():print(g,'S4',','.join(z['s4_memberships']),'S7 glucose',z['s7']['glucose'] if z['s7'] else 'MISSING')
